"""Build a short, editable research update for a nontechnical audience.

The deck deliberately separates the completed historical audit, authored
calibration cases, excluded pilots, and the planned prospective experiment.
Run with the Codex bundled Python, which includes python-pptx 1.0.2.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parent / "Research_Update_2026-11-01.pptx"

NAVY = "0B1220"
PANEL = "172238"
PANEL_2 = "1D2B45"
WHITE = "F6F4EF"
MUTED = "A8B5C8"
MINT = "5DDFC4"
CORAL = "FF7773"
BLUE = "77A9FF"
AMBER = "FFC96E"
LINE = "34445D"
FONT = "Arial"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def text(
    slide, value: str, x: float, y: float, w: float, h: float,
    *, size: float = 20, color: str = WHITE, bold: bool = False,
    align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE,
    margin: float = 0, line_spacing: float = 1.1,
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for idx, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)


def box(slide, x: float, y: float, w: float, h: float, *,
        fill: str = PANEL, border: str | None = None, round_: bool = True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if border:
        s.line.color.rgb = rgb(border)
        s.line.width = Pt(1.2)
    else:
        s.line.fill.background()
    if round_:
        s.adjustments[0] = 0.12
    return s


def circle(slide, x: float, y: float, d: float, *, fill: str,
           border: str | None = None):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if border:
        s.line.color.rgb = rgb(border)
        s.line.width = Pt(1.1)
    else:
        s.line.fill.background()
    return s


def rule(slide, x1: float, y1: float, x2: float, y2: float,
         *, color: str = LINE, width: float = 1.5):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    s.line.color.rgb = rgb(color)
    s.line.width = Pt(width)
    return s


def pill(slide, value: str, x: float, y: float, w: float, h: float,
         *, fill: str = PANEL_2, color: str = WHITE,
         size: float = 13, border: str | None = None) -> None:
    box(slide, x, y, w, h, fill=fill, border=border)
    text(slide, value, x + 0.10, y + 0.02, w - 0.20, h - 0.04,
         size=size, color=color, bold=True, align=PP_ALIGN.CENTER)


def base(prs: Presentation, number: int, eyebrow: str, title: str,
         source: str, *, title_size: float = 33):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box(slide, 0, 0, 13.333, 7.5, fill=NAVY, round_=False)
    text(slide, eyebrow.upper(), 0.70, 0.36, 10.9, 0.30,
         size=11, color=BLUE, bold=True)
    text(slide, title, 0.70, 0.81, 11.95, 0.92,
         size=title_size, color=WHITE, bold=True,
         valign=MSO_ANCHOR.TOP)
    rule(slide, 0.70, 7.08, 12.63, 7.08)
    text(slide, source, 0.72, 7.14, 11.4, 0.19,
         size=8.5, color=MUTED)
    text(slide, f"{number:02d}", 12.1, 7.11, 0.5, 0.24,
         size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    return slide


def speaker_note(slide, value: str) -> None:
    slide.notes_slide.notes_text_frame.text = value


def build() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "The Tests Passed. The Benchmark Said No."
    prs.core_properties.subject = "Research progress on AI security-repair evaluation"
    prs.core_properties.author = "Imran Zakirov"
    prs.core_properties.keywords = "AI coding agent, security repair, benchmark validity"

    # 1 — Hook.
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, 13.333, 7.5, fill=NAVY, round_=False)
    pill(s, "RESEARCH UPDATE  ·  1 NOV 2026", 0.72, 0.53, 3.63, 0.42,
         fill=PANEL_2, color=BLUE, size=11)
    text(s, "The Tests Passed.\nThe Benchmark Said No.", 0.70, 1.28, 10.9, 2.05,
         size=45, bold=True, valign=MSO_ANCHOR.TOP, line_spacing=1.03)
    text(s, "What if the grader of an AI security fix is wrong?", 0.73, 3.54,
         10.6, 0.55, size=23, color=MUTED)
    pill(s, "TESTS  PASS", 0.77, 4.62, 3.18, 0.72,
         fill="17453E", color=MINT, size=23, border="2D7769")
    text(s, "→", 4.23, 4.59, 0.75, 0.74, size=36, color=MUTED,
         align=PP_ALIGN.CENTER)
    pill(s, "FINAL VERDICT  FAIL", 5.20, 4.62, 4.48, 0.72,
         fill="4A2937", color=CORAL, size=22, border="875063")
    rule(s, 0.72, 6.78, 12.61, 6.78)
    text(s, "Imran Zakirov  ·  Bishkek, Kyrgyzstan", 0.73, 6.91, 8, 0.25,
         size=11, color=MUTED)
    text(s, "01", 12.08, 6.89, 0.5, 0.26, size=11, color=MUTED,
         align=PP_ALIGN.RIGHT)
    speaker_note(s, "We asked whether an AI that repairs code is being judged fairly. "
                 "Our headline is a verified disagreement inside one historical benchmark, "
                 "not a claim that the generated code was secure in general. "
                 "Source: README.md; paper/final.md, Results.")

    # 2 — Everyday analogy.
    s = base(prs, 2, "The idea", "Imagine a teacher with a checklist",
             "Illustration of the benchmark disagreement; see paper/final.md, Results.")
    box(s, 0.76, 2.07, 5.32, 3.56, fill=PANEL)
    text(s, "THE CHECKLIST", 1.09, 2.39, 4.45, 0.37,
         size=14, color=MUTED, bold=True)
    circle(s, 1.13, 3.08, 0.51, fill=MINT)
    text(s, "✓", 1.19, 3.05, 0.38, 0.54, size=28, color=NAVY,
         bold=True, align=PP_ALIGN.CENTER)
    text(s, "All listed checks pass", 1.82, 3.09, 3.85, 0.54,
         size=22, bold=True)
    text(s, "The student's work satisfies\nwhat the checklist actually tests.",
         1.13, 4.10, 4.38, 0.95, size=17, color=MUTED)
    text(s, "→", 6.27, 3.33, 0.8, 0.85,
         size=38, color=BLUE, align=PP_ALIGN.CENTER)
    box(s, 7.25, 2.07, 5.30, 3.56, fill="3D2634", border="794458")
    text(s, "THE FINAL GRADE", 7.60, 2.39, 4.45, 0.37,
         size=14, color="F7B0AC", bold=True)
    circle(s, 7.64, 3.08, 0.51, fill=CORAL)
    text(s, "×", 7.70, 3.05, 0.38, 0.54, size=30, color=NAVY,
         bold=True, align=PP_ALIGN.CENTER)
    text(s, "Still marked FAIL", 8.35, 3.09, 3.75, 0.54,
         size=22, bold=True)
    text(s, "An extra rule makes the final\ngrade disagree with the checks.",
         7.64, 4.10, 4.35, 0.95, size=17, color="E7B7BC")
    pill(s, "A checklist can also miss important problems.", 2.28, 6.02,
         8.78, 0.46, fill=PANEL_2, color=AMBER, size=14)
    speaker_note(s, "Think of a teacher whose own checklist is satisfied, yet who "
                 "marks the work wrong because of an additional rule. That is the "
                 "analogy for our historical checker. The checklist is limited; "
                 "passing it is not proof that the answer is fully correct.")

    # 3 — Agent feedback loop.
    s = base(prs, 3, "How it works", "An AI coder responds to the checker",
             "Study process: paper/final.md, Methods; README.md.")
    stages = [
        (0.76, "1", "A security task", "A small piece of code\nneeds repair.", BLUE),
        (3.89, "2", "AI proposes code", "The agent submits\na candidate fix.", MINT),
        (7.02, "3", "Checker responds", "Tests and final rule\nproduce feedback.", AMBER),
        (10.15, "4", "AI acts again", "It may revise, stop,\nor run out of time.", CORAL),
    ]
    for x, num, heading, body, accent in stages:
        box(s, x, 2.19, 2.42, 3.24, fill=PANEL)
        circle(s, x + 0.85, 2.51, 0.70, fill=accent)
        text(s, num, x + 0.91, 2.55, 0.58, 0.58, size=23,
             color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        text(s, heading, x + 0.17, 3.45, 2.08, 0.55,
             size=17.5, bold=True, align=PP_ALIGN.CENTER)
        text(s, body, x + 0.17, 4.18, 2.08, 0.79,
             size=13.5, color=MUTED, align=PP_ALIGN.CENTER)
    for x in (3.38, 6.51, 9.64):
        text(s, "→", x, 3.24, 0.36, 0.65, size=27,
             color=BLUE, align=PP_ALIGN.CENTER)
    text(s, "Our research asks: what can we infer from the AI's reaction if the feedback is wrong?",
         1.06, 5.92, 11.21, 0.68, size=22,
         color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    speaker_note(s, "The AI agent writes small Python repairs, receives feedback, "
                 "and can try again. We archive submissions and observations. "
                 "A final pass rate alone cannot tell us what happened along the way.")

    # 4 — Historical cohort.
    s = base(prs, 4, "What we examined", "One recorded study, clearly bounded",
             "Historical exploratory cohort: README.md, Current result; paper/final.md, Methods.")
    metrics = [
        (0.77, "40", "recorded repair\nepisodes", BLUE),
        (4.54, "5", "selected security\ntask families", MINT),
        (8.31, "1", "local coding\nmodel", AMBER),
    ]
    for x, numeral, label, accent in metrics:
        box(s, x, 2.01, 3.18, 2.56, fill=PANEL)
        text(s, numeral, x + 0.24, 2.27, 2.7, 0.98,
             size=55, color=accent, bold=True, align=PP_ALIGN.CENTER)
        text(s, label, x + 0.26, 3.43, 2.66, 0.78,
             size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(s, "The five families covered key derivation, nonce handling, weak randomness,",
         0.78, 5.08, 11.8, 0.38, size=19, color=MUTED)
    text(s, "SQL injection, and cross-site scripting.",
         0.78, 5.48, 11.8, 0.38, size=19, color=MUTED)
    pill(s, "HISTORICAL EXPLORATORY DATA", 0.78, 6.18, 4.12, 0.45,
         fill=PANEL_2, color=BLUE, size=12)
    speaker_note(s, "We audited 40 archived episodes from one local model, "
                 "qwen3:4b, across five selected security task families. One episode "
                 "is a recorded attempt under a particular condition. This is an "
                 "exploratory cohort, not the future main experiment.")

    # 5 — Key finding, dot plot.
    s = base(prs, 5, "The finding", "Tests passed. The final answer was still FAIL.",
             "Counts: docs/MEASUREMENT_FOLLOWUP_20261030.md, §1; paper/final.md, Results.",
             title_size=30)
    box(s, 0.78, 1.98, 7.20, 3.73, fill=PANEL)
    for i in range(20):
        row, col = divmod(i, 10)
        d = 0.43
        x = 1.26 + col * 0.61
        y = 2.68 + row * 0.84
        circle(s, x, y, d, fill=CORAL if i < 9 else PANEL_2,
               border=None if i < 9 else LINE)
    text(s, "20 final submissions in the nominally\nunavailable condition",
         1.24, 4.64, 6.19, 0.79, size=17, color=MUTED)
    text(s, "9 / 20", 8.43, 2.38, 3.96, 1.25,
         size=58, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
    text(s, "passed the implemented\nfamily tests, yet the\nfinal gate rejected them",
         8.40, 3.85, 4.00, 1.48, size=21,
         color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    pill(s, "This is a checker disagreement in selected runs.",
         1.72, 6.06, 9.90, 0.48, fill="3B2834", color="FFC0BC", size=15)
    speaker_note(s, "In nine of the 20 historical runs labeled nominally "
                 "unavailable, the final submission passed the implemented family "
                 "tests but the benchmark rejected it. That is nine of those "
                 "20 selected runs, not a general error rate. The receipt audit "
                 "reproduced these counts from the saved archive.")

    # 6 — Hidden gate mechanism.
    s = base(prs, 6, "Why it happened", "The benchmark had a second gate",
             "Evaluator rule and receipts: paper/final.md, Results; README.md.")
    box(s, 0.84, 2.17, 4.20, 1.82, fill="173D3B", border="2A716C")
    text(s, "IMPLEMENTED TESTS", 1.10, 2.47, 3.70, 0.38,
         size=14, color=MINT, bold=True, align=PP_ALIGN.CENTER)
    text(s, "PASS", 1.11, 3.04, 3.69, 0.56,
         size=29, color=MINT, bold=True, align=PP_ALIGN.CENTER)
    text(s, "AND", 5.23, 2.81, 1.04, 0.72,
         size=19, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    box(s, 6.42, 2.17, 4.20, 1.82, fill="3C2836", border="755065")
    text(s, "HIDDEN CONDITION FLAG", 6.68, 2.47, 3.70, 0.38,
         size=14, color="FFC0BC", bold=True, align=PP_ALIGN.CENTER)
    text(s, "FALSE", 6.69, 3.04, 3.69, 0.56,
         size=29, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
    text(s, "=", 10.85, 2.81, 0.52, 0.72,
         size=31, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    box(s, 11.46, 2.60, 1.05, 0.90, fill="633243", border="9F5469")
    text(s, "FAIL", 11.52, 2.72, 0.92, 0.56,
         size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(s, "The extra flag was false in the nominally unavailable condition.",
         0.86, 4.66, 11.73, 0.50, size=21, bold=True,
         align=PP_ALIGN.CENTER)
    text(s, "Passing these finite tests still does not prove that code is secure in general.",
         0.86, 5.44, 11.73, 0.55, size=18,
         color=AMBER, align=PP_ALIGN.CENTER)
    speaker_note(s, "The final rule was effectively: accept if the implemented "
                 "tests pass and a hidden condition flag allows acceptance. In the "
                 "nominally unavailable condition that flag was false. This "
                 "explains the test-versus-gate disagreement. It does not show "
                 "that the code was secure beyond those finite tests.")

    # 7 — Accomplishments, with separate cohorts.
    s = base(prs, 7, "What we built", "We checked the evidence from several angles",
             "README.md; docs/MEASUREMENT_FOLLOWUP_20261030.md; docs/V7_DECISIVE_EXCLUDED_PILOT_20261005.md.",
             title_size=31)
    cards = [
        (0.78, 2.05, BLUE, "ARCHIVE REPLAY", "Recomputed the 40-run result\nfrom saved receipts."),
        (6.72, 2.05, MINT, "TASK CATALOG", "Authored 20 finite task families;\nhuman review is still pending."),
        (0.78, 4.22, AMBER, "TEST CALIBRATION", "Rechecked 1,820 authored\ncandidate/condition decisions."),
        (6.72, 4.22, CORAL, "CORRECTED PILOT", "Completed 30 separate, excluded\nfeasibility episodes."),
    ]
    for x, y, accent, head, body in cards:
        box(s, x, y, 5.78, 1.77, fill=PANEL)
        box(s, x, y, 0.12, 1.77, fill=accent, round_=False)
        text(s, head, x + 0.31, y + 0.18, 4.98, 0.36,
             size=13, color=accent, bold=True)
        text(s, body, x + 0.30, y + 0.63, 5.05, 0.83,
             size=18.3, bold=True)
    speaker_note(s, "We rebuilt the historical result from saved records and "
                 "added automated verification in the public repository. "
                 "Separately, we authored a 20-family finite catalog and "
                 "independently rechecked 1,820 authored decisions. Those are "
                 "designed examples, not AI submissions. A corrected-oracle pilot "
                 "completed 30 excluded episodes in two authored families. "
                 "None of these cohorts is pooled with the 40 historical runs.")

    # 8 — Limits and scientific discipline.
    s = base(prs, 8, "What the evidence means", "A clear result, with clear limits",
             "Limitations: paper/final.md, Results and Discussion; README.md.")
    box(s, 0.78, 2.04, 5.84, 3.49, fill="173A39", border="2A716C")
    text(s, "WE CAN SAY", 1.08, 2.34, 5.10, 0.42,
         size=15, color=MINT, bold=True)
    text(s, "The checker and its own\nimplemented tests disagreed\nin nine archived finals.",
         1.08, 3.00, 5.06, 1.91, size=24, bold=True,
         valign=MSO_ANCHOR.TOP)
    box(s, 6.76, 2.04, 5.78, 3.49, fill="3A2C35", border="795867")
    text(s, "WE CANNOT YET SAY", 7.05, 2.34, 5.10, 0.42,
         size=15, color=CORAL, bold=True)
    text(s, "The AI generally reacts badly\nto wrong feedback—or that\nthose fixes were fully secure.",
         7.06, 3.00, 5.06, 1.91, size=23, bold=True,
         valign=MSO_ANCHOR.TOP)
    pill(s, "Old intermediate logs: 0 of 81 action receipts stored raw test results.",
         0.80, 5.96, 11.73, 0.52,
         fill=PANEL_2, color=AMBER, size=15)
    speaker_note(s, "We can say the historical checker was internally "
                 "inconsistent in nine final submissions. We cannot say this "
                 "caused an agent to continue or that the code was truly secure. "
                 "The old per-action receipts did not preserve raw test results "
                 "for any of the 81 intermediate checks, so we cannot identify "
                 "every point where a passing intermediate candidate appeared.")

    # 9 — Prospective causal design.
    s = base(prs, 9, "The next experiment", "Same code. Two different verdicts.",
             "Planned design only: docs/INTEGRATED_STUDY_EXTENSION_20261018.md.")
    box(s, 4.15, 1.95, 5.04, 1.04, fill=PANEL_2, border=LINE)
    text(s, "One independently checked candidate", 4.39, 2.18, 4.55, 0.52,
         size=19.5, bold=True, align=PP_ALIGN.CENTER)
    rule(s, 6.67, 2.99, 6.67, 3.44, color=BLUE, width=2)
    rule(s, 3.70, 3.44, 9.63, 3.44, color=BLUE, width=2)
    rule(s, 3.70, 3.44, 3.70, 3.77, color=BLUE, width=2)
    rule(s, 9.63, 3.44, 9.63, 3.77, color=BLUE, width=2)
    box(s, 1.20, 3.77, 5.04, 1.30, fill="173D3B", border="2A716C")
    text(s, "A  ·  truthful PASS", 1.51, 4.11, 4.40, 0.59,
         size=23, color=MINT, bold=True, align=PP_ALIGN.CENTER)
    box(s, 7.08, 3.77, 5.04, 1.30, fill="3C2836", border="755065")
    text(s, "B  ·  false FAIL", 7.40, 4.11, 4.40, 0.59,
         size=23, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
    text(s, "Compare the later code and whether it keeps the verified finite property.",
         1.02, 5.64, 11.27, 0.55, size=20,
         bold=True, align=PP_ALIGN.CENTER)
    pill(s, "PLANNED  ·  NO ELIGIBLE MAIN RESULTS YET", 3.74, 6.29,
         5.84, 0.42, fill=PANEL_2, color=BLUE, size=12)
    speaker_note(s, "The planned study branches from the same agent-generated, "
                 "independently verified candidate and the same previous transcript. "
                 "One branch hears a truthful PASS; the other a false FAIL. "
                 "We would measure early stopping and whether the final candidate "
                 "retains the bounded security property. A selected excluded "
                 "pilot pair motivated this question, but it cannot establish "
                 "a general causal effect. The main study has no eligible data.")

    # 10 — Takeaway + current status.
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, 13.333, 7.5, fill=NAVY, round_=False)
    text(s, "THE TAKEAWAY", 0.73, 0.54, 9.0, 0.35,
         size=12, color=BLUE, bold=True)
    text(s, "Before judging the AI,\ncheck the judge.",
         0.73, 1.27, 11.62, 1.77, size=48, bold=True,
         valign=MSO_ANCHOR.TOP, line_spacing=1.03)
    box(s, 0.77, 3.74, 11.78, 1.39, fill=PANEL)
    text(s, "COMPLETED", 1.05, 4.00, 2.05, 0.30,
         size=12, color=MINT, bold=True)
    text(s, "historical audit", 1.06, 4.38, 2.82, 0.42,
         size=20, bold=True)
    text(s, "NEXT", 4.60, 4.00, 1.15, 0.30,
         size=12, color=AMBER, bold=True)
    text(s, "independent review", 4.60, 4.38, 3.01, 0.42,
         size=20, bold=True)
    text(s, "MAIN STUDY", 8.80, 4.00, 2.12, 0.30,
         size=12, color=CORAL, bold=True)
    text(s, "0 / 12 gates passed", 8.81, 4.38, 3.22, 0.42,
         size=20, bold=True)
    text(s, "as of 1 Nov 2026", 8.82, 4.83, 3.12, 0.20,
         size=10, color=MUTED)
    text(s, "Would you trust an AI score if the grader disagreed with its own tests?",
         0.73, 5.60, 11.90, 0.59, size=21, color=MUTED)
    text(s, "github.com/Imranqsl212/AiResearch", 0.74, 6.65,
         8.05, 0.35, size=16, color=BLUE, bold=True)
    rule(s, 0.73, 7.13, 12.62, 7.13)
    text(s, "Sources: README.md · paper/final.md · docs/MEASUREMENT_FOLLOWUP_20261030.md",
         0.74, 7.18, 11.5, 0.18, size=8.5, color=MUTED)
    text(s, "10", 12.10, 7.17, 0.5, 0.2,
         size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    speaker_note(s, "Our completed finding is about measurement integrity. "
                 "The future causal question is exciting, and it requires "
                 "independently reviewed tasks and a frozen main protocol. "
                 "As of this presentation, the prospective release gate is "
                 "0 of 12 and there are zero eligible main episodes. "
                 "The full repository is public at github.com/Imranqsl212/AiResearch.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
