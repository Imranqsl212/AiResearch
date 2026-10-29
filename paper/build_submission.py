"""Rebuild publication figures and PDF/DOCX from derived, auditable data.

Run ``python3 -B -m analysis.paper_pipeline`` first, then
``python3 -B -m paper.build_submission`` with the workspace Python runtime.
This renderer never loads or executes candidate code or calls an agent.
"""

from __future__ import annotations

import csv
import html
import io
import json
import re
from copy import copy
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as PdfImage
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle


ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper"
FIGURES = PAPER / "figures"
SUBMISSION = PAPER / "submission"
ARTIFACTS = PAPER / "artifacts"
FONT_DIR = Path("/System/Library/Fonts/Supplemental")
FONT_REG = FONT_DIR / "Arial.ttf"
FONT_BOLD = FONT_DIR / "Arial Bold.ttf"
NAVY = "#16324F"
TEAL = "#1F8A8A"
ORANGE = "#C36E39"
GRAY = "#687786"
LIGHT = "#E9EFF3"
INK = "#17212B"
AUTHOR = "Imran Zakirov"
PAGE_LIMIT = 20
BODY_POINT_SIZE = 12


def review_manuscript(markdown: str) -> str:
    """Remove identity-bearing front matter and acknowledgments for review.

    The scientific prose, results, tables, figures, and references come from the
    same source for both versions. Fail closed if a known identifier remains.
    """
    kept: list[str] = []
    in_acknowledgments = False
    for line in markdown.splitlines():
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if heading and len(heading.group(1)) <= 2:
            in_acknowledgments = bool(re.fullmatch(
                r"Acknowledg(?:e)?ments", heading.group(2), re.IGNORECASE))
            if in_acknowledgments:
                kept.extend((line, "", "Acknowledgments and author disclosures are withheld "
                             "during anonymous peer review.", ""))
                continue
        if in_acknowledgments:
            continue
        if re.match(r"^\*\*(?:Author(?: line)?|Authors|Affiliation(?:s)?|Correspondence|Corresponding author):\*\*",
                    line, re.IGNORECASE):
            continue
        # The study commit may expose the author through its repository. Keep
        # the substantive data-access limitation without a linkable revision.
        if line.startswith("**Data and code:**"):
            line = ("**Data and code:** Source trajectories are tracked in the "
                    "research branch; external access depends on the exact public "
                    "revision. Derived tables and analysis code accompany the manuscript.")
        line = line.replace("The frozen study commit is given above; paper analysis code is a post collection addition.",
                            "Paper analysis code is a post collection addition.")
        kept.append(line)
    review = "\n".join(kept).rstrip() + "\n"
    review = review.replace(
        "The original study source commit is `98a4dcf414672c1725e4acdb9f08045370aa7f64`;",
        "The original study source revision is withheld during anonymous review;",
    )
    review = review.replace("98a4dcf414672c1725e4acdb9f08045370aa7f64",
                            "source revision withheld during anonymous review")
    review = review.replace(
        "The public repository's data revision `37074c198237ade67eb76921333b2835d329c863` "
        "contains the raw trajectories and archives needed to recalculate the reported counts;",
        "A source revision withheld during anonymous review contains the raw trajectories "
        "and archives needed to recalculate the reported counts;",
    )
    if re.search(
        r"(?i)\b(?:imran|zakirov)\b|zakirovimran333@gmail\.com|"
        r"37074c198237ade67eb76921333b2835d329c863|github\.com/Imranqsl212",
        review,
    ):
        raise ValueError("anonymized manuscript still contains a known author identifier")
    return review


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = FONT_BOLD if bold and FONT_BOLD.is_file() else FONT_REG
    return ImageFont.truetype(str(path), size)


def label(draw: ImageDraw.ImageDraw, xy: tuple[int, int], value: str,
          *, size: int = 32, bold: bool = False, fill: str = INK,
          anchor: str | None = None) -> None:
    draw.text(xy, str(value), font=font(size, bold), fill=fill, anchor=anchor)


def canvas(title: str, subtitle: str, *, width: int = 1600, height: int = 900):
    im = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(im)
    draw.rectangle((0, 0, width, 14), fill=NAVY)
    label(draw, (105, 75), title, size=46, bold=True)
    label(draw, (105, 137), subtitle, size=25, fill=GRAY)
    draw.line((105, height - 80, width - 105, height - 80), fill=LIGHT, width=2)
    label(draw, (105, height - 56), "Exploratory local study · source: derived run metrics",
          size=19, fill=GRAY)
    return im, draw


def save(im: Image.Image, number: int, stem: str) -> None:
    FIGURES.mkdir(parents=True, exist_ok=True)
    im.save(FIGURES / f"figure{number}_{stem}.png", optimize=True, dpi=(220, 220))


def figure1(rows: list[dict], results: dict) -> None:
    im, d = canvas("Finite-check acceptance by family", "Repairable cells only · 2 episodes per bar · n=20 main episodes")
    families = sorted({r["family"] for r in rows if r["feasibility"] == "R"})
    plotted = [(family, mode, sum(r["accepted_final_candidate"] == "True" for r in rows
                                   if r["family"] == family and r["feasibility"] == "R" and r["retrieval"] == mode))
               for family in families for mode in ("off", "relevant")]
    y_top, y_bottom = 270, 660
    for tick in range(3):
        y = y_bottom - tick * (y_bottom - y_top) // 2
        d.line((160, y, 1480, y), fill=LIGHT, width=2)
        label(d, (120, y), str(tick), size=27, anchor="rm")
    colors_map = {"off": NAVY, "relevant": TEAL}
    for i, family in enumerate(families):
        cx = 280 + i * 260
        for j, mode in enumerate(("off", "relevant")):
            value = next(v for f, m, v in plotted if f == family and m == mode)
            x1 = cx + (j * 72) - 56
            d.rectangle((x1, y_bottom - value * 195, x1 + 60, y_bottom), fill=colors_map[mode])
            label(d, (x1 + 30, y_bottom - value * 195 - 13), str(value), size=29, bold=True, anchor="mb")
        shown = family.replace("-", " ")
        label(d, (cx - 22, y_bottom + 29), shown, size=25, anchor="mt")
    d.rectangle((590, 755, 620, 785), fill=NAVY); label(d, (638, 770), "Guidance off", size=25, anchor="lm")
    d.rectangle((890, 755, 920, 785), fill=TEAL); label(d, (938, 770), "Fixed guidance", size=25, anchor="lm")
    label(d, (105, 745), "Count passing", size=25)
    save(im, 1, "family_acceptance")


def figure2(rows: list[dict]) -> None:
    im, d = canvas("Tool actions per episode", "Observed tool calls per episode · R n=20; U n=20")
    counts = {(group, k): sum(r["feasibility"] == group and int(r["actions"]) == k for r in rows)
              for group in ("R", "U") for k in range(2, 7)}
    # Scale from the observed maximum so a tall bar cannot enter the title.
    upper = max(16, 4 * ((max(counts.values()) + 3) // 4))
    pixels_per_episode = 420 / upper
    for tick in range(0, upper + 1, upper // 4):
        y = 690 - tick * pixels_per_episode
        d.line((175, y, 1470, y), fill=LIGHT, width=2)
        label(d, (140, y), str(tick), size=26, anchor="rm")
    for i, step in enumerate(range(2, 7)):
        cx = 305 + i * 260
        for j, group in enumerate(("R", "U")):
            count = counts[(group, step)]
            x1 = cx + j * 82 - 78
            d.rectangle((x1, 690 - count * pixels_per_episode, x1 + 72, 690),
                        fill=NAVY if group == "R" else ORANGE)
            label(d, (x1 + 36, 680 - count * pixels_per_episode), str(count),
                  size=26, bold=True, anchor="mb")
        label(d, (cx, 715), str(step), size=31, anchor="mt")
    label(d, (950, 755), "Tool actions per episode", size=25, anchor="mm")
    label(d, (105, 750), "Episodes", size=25)
    d.rectangle((570, 773, 600, 803), fill=NAVY); label(d, (610, 788), "R", size=25, anchor="lm")
    d.rectangle((690, 773, 720, 803), fill=ORANGE); label(d, (730, 788), "U", size=25, anchor="lm")
    save(im, 2, "actions")


def figure3(results: dict) -> None:
    im, d = canvas("How much source change is measurable?", "Main n=40 episodes · 77 submitted candidates · source pairs have nested denominators")
    values = [
        ("Adjacent source pairs", results["main"]["source_pairs"], "37 / 37 observed pairs"),
        ("Both sources parseable", results["main"]["parseable_source_pairs"], "25 / 37 pairs"),
        ("Called API set changed", results["main"]["api_inventory_change_pairs"], "17 / 25 parseable pairs"),
    ]
    for tick in range(0, 41, 10):
        x = 610 + tick * 20
        d.line((x, 275, x, 680), fill=LIGHT, width=2)
        label(d, (x, 698), str(tick), size=25, anchor="mt")
    for i, (name, value, denominator) in enumerate(values):
        y = 320 + i * 145
        label(d, (550, y + 35), name, size=30, anchor="rm")
        d.rectangle((610, y, 610 + value * 20, y + 70), fill=[NAVY, TEAL, ORANGE][i])
        label(d, (625 + value * 20, y + 35), denominator, size=24, anchor="lm")
    label(d, (1000, 746), "Count of adjacent pairs", size=26, anchor="mm")
    save(im, 3, "source_pairs")


def figure4(rows: list[dict]) -> None:
    im, d = canvas("Recorded stop events", "One terminal event per main episode · R n=20; U n=20 · event label only")
    categories = [("AGENT_SELF_TERMINATION", NAVY, "self termination"),
                  ("BUDGET_STOP", ORANGE, "budget"), ("TIMEOUT", TEAL, "timeout")]
    for i, group in enumerate(("R", "U")):
        y = 350 + i * 180
        label(d, (170, y + 35), group, size=42, bold=True, anchor="rm")
        left = 230
        for category, color, _ in categories:
            n = sum(r["feasibility"] == group and r["stop_event"] == category for r in rows)
            if n:
                d.rectangle((left, y, left + n * 60, y + 85), fill=color)
                label(d, (left + n * 30, y + 42), str(n), size=30, bold=True, fill="white", anchor="mm")
                left += n * 60
        label(d, (1460, y + 45), "n=20", size=24, anchor="lm")
    x = 230
    for _, color, name in categories:
        d.rectangle((x, 730, x + 25, 755), fill=color)
        label(d, (x + 34, 743), name, size=23, anchor="lm")
        x += 340
    label(d, (105, 690), "Episodes (each bar totals 20)", size=25)
    save(im, 4, "stopping")


def figure5(rows: list[dict]) -> None:
    im, d = canvas("Explicit final public claims", "Conservative claim recovery · 40 main episodes · missingness shown")
    counts = Counter(r["recovered_final_claim"] for r in rows)
    for tick in range(0, 21, 5):
        y = 690 - tick * 21
        d.line((200, y, 1450, y), fill=LIGHT, width=2)
        label(d, (160, y), str(tick), size=25, anchor="rm")
    for i, (name, key, color) in enumerate([("Success", "success", NAVY),
                                           ("Non success", "non_success", TEAL),
                                           ("Unrecovered", "UNRECOVERED", GRAY)]):
        count = counts[key]
        x = 355 + i * 430
        d.rectangle((x, 690 - count * 21, x + 180, 690), fill=color)
        label(d, (x + 90, 678 - count * 21), str(count), size=38, bold=True, anchor="mb")
        label(d, (x + 90, 716), name, size=29, anchor="mt")
    label(d, (106, 745), "Episodes", size=25)
    save(im, 5, "claims")


def figure6(rows: list[dict]) -> None:
    im, d = canvas("Five observable trajectories", "Public tool observations by step · purposeful illustrations, not a random sample", width=1800, height=1100)
    slots = [1, 10, 15, 22, 39]
    by_slot = {r["slot_id"]: r for r in rows}
    colors_map = {"SECURITY_CHECK_PASSED": TEAL, "SECURITY_CHECK_FAILED": ORANGE}
    for i, slot in enumerate(slots):
        row = by_slot[f"slot-{slot:06d}"]
        y = 320 + i * 137
        d.line((630, y, 1620, y), fill=LIGHT, width=4)
        label(d, (110, y - 15), f"{slot:02d} · {row['family']} / {row['condition']}", size=30, bold=True)
        label(d, (112, y + 27), f"guidance {row['retrieval']}; {row['stop_event'].lower().replace('_', ' ')}", size=22, fill=GRAY)
        for item in row["sequence"].split(" | "):
            parts = item.split(":", 2)
            if len(parts) != 3:
                continue
            step, tool, outcome = parts
            x = 660 + (int(step) - 1) * 190
            color = colors_map.get(outcome, NAVY)
            d.ellipse((x - 14, y - 14, x + 14, y + 14), fill=color)
            label(d, (x, y - 46), f"{step}:{tool}", size=20, anchor="mm")
            label(d, (x, y + 28), "pass" if outcome == "SECURITY_CHECK_PASSED" else
                  "fail" if outcome == "SECURITY_CHECK_FAILED" else outcome.lower().replace("_", " "),
                  size=20, fill=color, anchor="mt")
    d.rectangle((535, 991, 560, 1016), fill=TEAL); label(d, (570, 1004), "public check passed", size=24, anchor="lm")
    d.rectangle((970, 991, 995, 1016), fill=ORANGE); label(d, (1005, 1004), "public check failed", size=24, anchor="lm")
    d.rectangle((1335, 991, 1360, 1016), fill=NAVY); label(d, (1370, 1004), "other tool observation", size=24, anchor="lm")
    save(im, 6, "timelines")


def read_rows() -> tuple[list[dict], dict]:
    with (ARTIFACTS / "run_metrics.csv").open(encoding="utf-8", newline="") as handle:
        rows = [row for row in csv.DictReader(handle) if row["phase"] == "main"]
    results = json.loads((ARTIFACTS / "results.json").read_text(encoding="utf-8"))
    # Refuse to typeset a manuscript with stale or incorrectly pooled analysis.
    checks = [len(rows) == 40, results["main"]["accepted"] == 16,
              results["repairable"]["accepted"] == 16,
              results["main"]["test_gate_discordance"] == 9,
              results["main"]["check_observations_total"] == 81,
              results["main"]["action_raw_check_receipts_total"] == 0,
              results["policy_rejected"]["hidden_policy_veto"] == 9,
              results["pilot"]["n"] == 8,
              results["main"]["claim_counts"] == {"unknown": 40},
              results["confirmatory_GEE"] == "NOT ESTIMABLE / NOT RUN"]
    if not all(checks):
        raise ValueError("derived data do not match manuscript's verified cohort anchors")
    return rows, results


def inline(text: str) -> str:
    value = html.escape(text)
    value = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r'<link href="\2" color="#1F627F">\1</link>', value)
    value = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", value)
    value = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<i>\1</i>", value)
    value = re.sub(r"`([^`]+)`", r'<font name="Body">\1</font>', value)
    return value


def markdown_blocks(markdown: str):
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("!["):
            match = re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)", line)
            if match is None:
                raise ValueError(f"invalid figure reference: {line}")
            yield "figure", match.groups()
            index += 1
            continue
        if line.startswith("#"):
            depth = len(line) - len(line.lstrip("#"))
            yield "heading", (depth, line[depth:].strip())
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [c.strip() for c in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                index += 1
            yield "table", rows
            continue
        if re.match(r"^\d+\.\s", line):
            yield "list", line
            index += 1
            continue
        paragraph = [line]
        index += 1
        while index < len(lines) and lines[index].strip() and not lines[index].lstrip().startswith(("#", "|", "![")):
            if re.match(r"^\d+\.\s", lines[index].strip()):
                break
            paragraph.append(lines[index].strip())
            index += 1
        yield "paragraph", " ".join(paragraph)


def make_pdf(markdown: str, destination: Path, *, author: str) -> int:
    pdfmetrics.registerFont(TTFont("Body", str(FONT_REG)))
    pdfmetrics.registerFont(TTFont("BodyBold", str(FONT_BOLD if FONT_BOLD.is_file() else FONT_REG)))
    pdfmetrics.registerFontFamily("Body", normal="Body", bold="BodyBold", italic="Body", boldItalic="BodyBold")
    styles = getSampleStyleSheet()
    width, _ = A4
    body = ParagraphStyle("BodyX", parent=styles["Normal"], fontName="Body", fontSize=BODY_POINT_SIZE,
                          leading=BODY_POINT_SIZE, spaceAfter=3, textColor=colors.HexColor(INK))
    head1 = ParagraphStyle("H1X", parent=body, fontName="BodyBold", fontSize=BODY_POINT_SIZE,
                           leading=BODY_POINT_SIZE, spaceBefore=7, spaceAfter=3, keepWithNext=True,
                           textColor=colors.HexColor(NAVY))
    head2 = ParagraphStyle("H2X", parent=head1, spaceBefore=5)
    head3 = ParagraphStyle("H3X", parent=head2, leftIndent=10,
                           spaceBefore=3, textColor=colors.HexColor(INK))
    title = ParagraphStyle("TitleX", parent=head1, alignment=TA_CENTER, spaceAfter=5)
    caption = ParagraphStyle("CaptionX", parent=body, spaceBefore=2, spaceAfter=5,
                             textColor=colors.HexColor(GRAY))
    table_caption = ParagraphStyle("TableCaptionX", parent=body, keepWithNext=True,
                                   spaceBefore=3, spaceAfter=3)
    story = []
    for kind, payload in markdown_blocks(markdown):
        if kind == "heading":
            depth, value = payload
            style = title if depth == 1 else head1 if depth == 2 else head2 if depth == 3 else head3
            story.append(Paragraph(inline(value), style))
        elif kind == "paragraph":
            story.append(Paragraph(inline(payload), table_caption if payload.startswith("**Table ") else body))
        elif kind == "list":
            story.append(Paragraph(inline(payload), ParagraphStyle("ListX", parent=body, leftIndent=16, firstLineIndent=-16)))
        elif kind == "table":
            widths = [7.1 * inch / len(payload[0])] * len(payload[0])
            cell_style = ParagraphStyle("CellX", parent=body, spaceAfter=0)
            cells = [[Paragraph(inline(cell), cell_style) for cell in row] for row in payload]
            table = LongTable(cells, colWidths=widths, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(LIGHT)),
                ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor(NAVY)),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F9FB")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.extend([table, Spacer(1, 4)])
        elif kind == "figure":
            figcaption, relative = payload
            image_path = PAPER / relative
            if not image_path.is_file():
                raise FileNotFoundError(image_path)
            with Image.open(image_path) as im:
                ratio = im.height / im.width
            figure_width = 6.7 * inch
            story.append(PdfImage(str(image_path), width=figure_width, height=figure_width * ratio))
            story.append(Paragraph(inline(figcaption), caption))
    title_match = re.search(r"^#\s+(.+)$", markdown, re.MULTILINE)
    if title_match is None:
        raise ValueError("manuscript title is missing")
    doc = SimpleDocTemplate(str(destination), pagesize=A4, rightMargin=0.52*inch,
                            leftMargin=0.52*inch, topMargin=0.52*inch, bottomMargin=0.52*inch,
                            title=title_match.group(1), author=author,
                            invariant=1)

    def footer(canv, document):
        canv.saveState()
        canv.setFont("Body", BODY_POINT_SIZE)
        canv.setFillColor(colors.HexColor(GRAY))
        canv.drawString(0.52*inch, 0.3*inch, "Exploratory study")
        canv.drawRightString(width - 0.52*inch, 0.3*inch, f"{document.page}")
        canv.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    from pypdf import PdfReader
    pdf = PdfReader(str(destination))
    page_text = [page.extract_text() or "" for page in pdf.pages]
    extracted = "\n".join(page_text)
    required_sections = ("Abstract", "Introduction", "Methods", "Results", "Discussion", "References")
    for required in (*required_sections, "16/20", "9/20"):
        if required not in extracted:
            raise ValueError(f"PDF text extraction missing {required}")
    for caption, header in (("Table 1.", "Decision stage"), ("Table 2.", "Measure"),
                            ("Table 3.", "Fixed guidance")):
        if not any(caption in page and header in page for page in page_text):
            raise ValueError(f"PDF separates {caption} from its table header")
    return len(pdf.pages)


def make_docx(markdown: str, destination: Path, *, author: str) -> None:
    doc = Document()
    title_match = re.search(r"^#\s+(.+)$", markdown, re.MULTILINE)
    if title_match is None:
        raise ValueError("manuscript title is missing")
    doc.core_properties.title = title_match.group(1)
    doc.core_properties.author = author
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""
    doc.core_properties.keywords = ""
    doc.core_properties.subject = ""
    section = doc.sections[0]
    section.page_width = Pt(A4[0]); section.page_height = Pt(A4[1])
    section.top_margin = Inches(0.52); section.bottom_margin = Inches(0.52)
    section.left_margin = Inches(0.52); section.right_margin = Inches(0.52)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(BODY_POINT_SIZE)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(3 if name == "Normal" else 5)
    for kind, payload in markdown_blocks(markdown):
        if kind == "heading":
            depth, value = payload
            doc.add_heading(value, level=0 if depth == 1 else min(depth - 1, 3))
        elif kind in ("paragraph", "list"):
            p = doc.add_paragraph(style="Normal")
            text = payload if kind == "paragraph" else payload
            # Keep visible, editable prose and reference URLs in Word; Markdown
            # emphasis markers are removed only for display.
            text = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", text)
            text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", text)
            p.add_run(text.replace("**", "").replace("`", ""))
            if kind == "paragraph" and payload.startswith("**Table "):
                p.paragraph_format.keep_with_next = True
            if kind == "list":
                p.paragraph_format.left_indent = Inches(0.25)
        elif kind == "table":
            table = doc.add_table(rows=1, cols=len(payload[0]))
            table.style = "Light Shading Accent 1"
            for j, value in enumerate(payload[0]):
                table.rows[0].cells[j].text = value.replace("`", "")
            for row in payload[1:]:
                cells = table.add_row().cells
                for j, value in enumerate(row):
                    cells[j].text = value.replace("`", "")
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.style = doc.styles["Normal"]
                        paragraph.paragraph_format.space_after = Pt(0)
                        paragraph.paragraph_format.line_spacing = 1.0
        elif kind == "figure":
            caption_text, relative = payload
            doc.add_picture(str(PAPER / relative), width=Inches(6.7))
            p = doc.add_paragraph(caption_text)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.italic = True; run.font.size = Pt(BODY_POINT_SIZE)
    doc.save(destination)
    # python-docx stamps every ZIP member with the wall clock. Keep the OOXML
    # payloads and member metadata, but fix their ZIP timestamps for a stable
    # artifact hash when the manuscript, figures, and renderer are unchanged.
    normalized = io.BytesIO()
    with ZipFile(destination, "r") as source, ZipFile(normalized, "w") as target:
        for member in source.infolist():
            fixed = copy(member)
            fixed.date_time = (1980, 1, 1, 0, 0, 0)
            target.writestr(fixed, source.read(member.filename))
    destination.write_bytes(normalized.getvalue())


def verify_anonymous_artifacts(pdf_path: Path, docx_path: Path) -> None:
    """Check rendered text, package content, and document properties for identity."""
    from pypdf import PdfReader

    reader = PdfReader(str(pdf_path))
    pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    pdf_properties = "\n".join(str(value) for value in (reader.metadata or {}).values())
    with ZipFile(docx_path) as archive:
        docx_xml = "\n".join(
            archive.read(name).decode("utf-8", errors="replace")
            for name in archive.namelist() if name.endswith((".xml", ".rels"))
        )
    identity = re.compile(r"(?i)\b(?:imran|zakirov|bishkek)\b|zakirovimran333@gmail\.com|"
                          r"98a4dcf414672c1725e4acdb9f08045370aa7f64|"
                          r"37074c198237ade67eb76921333b2835d329c863|"
                          r"github\.com/Imranqsl212")
    for label, content in (("PDF text", pdf_text), ("PDF metadata", pdf_properties),
                           ("DOCX XML and metadata", docx_xml)):
        if identity.search(content):
            raise ValueError(f"anonymous {label} contains an author identifier")
    if reader.metadata and str(reader.metadata.get("/Author", "")).strip():
        raise ValueError("anonymous PDF author metadata is not blank")
    docx = Document(docx_path)
    if docx.core_properties.author or docx.core_properties.last_modified_by:
        raise ValueError("anonymous DOCX author metadata is not blank")


def main() -> None:
    rows, results = read_rows()
    figure1(rows, results); figure2(rows); figure3(results)
    figure4(rows); figure5(rows); figure6(rows)
    SUBMISSION.mkdir(parents=True, exist_ok=True)
    markdown = (PAPER / "final.md").read_text(encoding="utf-8")
    anonymous = review_manuscript(markdown)
    displays = Counter(kind for kind, _ in markdown_blocks(markdown)
                       if kind in ("figure", "table"))
    if displays.total() < 5:
        raise ValueError("research article requires at least five figures/tables")
    paths = {
        "identified_pdf": SUBMISSION / "full_paper.pdf",
        "identified_docx": SUBMISSION / "full_paper.docx",
        "anonymous_pdf": SUBMISSION / "full_paper_anonymous.pdf",
        "anonymous_docx": SUBMISSION / "full_paper_anonymous.docx",
    }
    identified_pages = make_pdf(markdown, paths["identified_pdf"], author=AUTHOR)
    make_docx(markdown, paths["identified_docx"], author=AUTHOR)
    anonymous_pages = make_pdf(anonymous, paths["anonymous_pdf"], author="")
    make_docx(anonymous, paths["anonymous_docx"], author="")
    verify_anonymous_artifacts(paths["anonymous_pdf"], paths["anonymous_docx"])
    pages = {"identified": identified_pages, "anonymous": anonymous_pages}
    within_limit = all(count <= PAGE_LIMIT for count in pages.values())
    print(json.dumps({"status": "BUILT" if within_limit else "PAGE_LIMIT_EXCEEDED",
                      "pdf_pages": pages, "page_limit": PAGE_LIMIT,
                      "body_font_pt": BODY_POINT_SIZE, "line_spacing": "single",
                      "figures": displays["figure"], "tables": displays["table"],
                      "main_episodes": len(rows),
                      "files": {key: str(path.relative_to(ROOT)) for key, path in paths.items()}}, indent=2))
    if not within_limit:
        raise ValueError(f"20-page limit exceeded: {pages}; content and 12-point font retained")


if __name__ == "__main__":
    main()
